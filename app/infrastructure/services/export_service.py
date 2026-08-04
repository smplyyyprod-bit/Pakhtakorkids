"""Excel and PDF export."""

import io
from datetime import date
from pathlib import Path

from sqlalchemy.ext.asyncio import AsyncSession

from app.core.logger import get_logger
from app.domain.models.daily_report import AttendanceStatus, UniformStatus
from app.infrastructure.repositories import CoachRepository, DailyReportRepository

logger = get_logger()

MONTHS_RU = {
    1: "Январь", 2: "Февраль", 3: "Март", 4: "Апрель",
    5: "Май", 6: "Июнь", 7: "Июль", 8: "Август",
    9: "Сентябрь", 10: "Октябрь", 11: "Ноябрь", 12: "Декабрь",
}

ATTENDANCE_RU = {
    AttendanceStatus.PRESENT: "Присутствует",
    AttendanceStatus.SICK_LEAVE: "Больничный",
    AttendanceStatus.VACATION: "Отпуск",
    AttendanceStatus.UNEXCUSED_ABSENCE: "Неуважительная",
    AttendanceStatus.NOT_FILLED: "Не заполнено",
}

UNIFORM_RU = {
    UniformStatus.YES: "Да",
    UniformStatus.NO: "Нет",
    UniformStatus.NO_DATA: "Нет данных",
}

# Candidate Unicode fonts. reportlab's built-in Helvetica has no Cyrillic
# glyphs at all - using it would silently render Russian text as blanks.
_FONT_CANDIDATES = [
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/usr/share/fonts/dejavu/DejaVuSans.ttf",
     "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf"),
    ("/usr/share/fonts/TTF/DejaVuSans.ttf",
     "/usr/share/fonts/TTF/DejaVuSans-Bold.ttf"),
]

FONT_REGULAR = "DejaVuSans"
FONT_BOLD = "DejaVuSans-Bold"

_fonts_ready: bool | None = None


def _register_fonts() -> bool:
    """Register a Cyrillic-capable font family exactly once.

    Returns False when no suitable font is installed, which the caller surfaces
    as a clear error rather than producing a PDF full of empty boxes.
    """
    global _fonts_ready
    if _fonts_ready is not None:
        return _fonts_ready

    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for regular, bold in _FONT_CANDIDATES:
        if Path(regular).exists() and Path(bold).exists():
            pdfmetrics.registerFont(TTFont(FONT_REGULAR, regular))
            pdfmetrics.registerFont(TTFont(FONT_BOLD, bold))
            from reportlab.lib.fonts import addMapping

            addMapping(FONT_REGULAR, 0, 0, FONT_REGULAR)
            addMapping(FONT_REGULAR, 1, 0, FONT_BOLD)
            _fonts_ready = True
            logger.info("Registered PDF font family from {}", regular)
            return True

    logger.error(
        "No Cyrillic TTF font found - install fonts-dejavu-core for PDF export"
    )
    _fonts_ready = False
    return False


class ExportService:
    def __init__(self, session: AsyncSession) -> None:
        self.session = session
        self.reports = DailyReportRepository(session)
        self.coaches = CoachRepository(session)

    async def _gather(self, coach_id: int, year: int, month: int):
        coach = await self.coaches.get_with_branch(coach_id)
        if coach is None:
            raise ValueError(f"Coach {coach_id} not found")
        rows = await self.reports.get_by_coach_and_month(coach_id, year, month)
        stats = await self.reports.monthly_stats(coach_id, year, month)
        return coach, rows, stats

    @staticmethod
    def _summary_rows(stats) -> list[tuple[str, object]]:
        return [
            ("Отработано часов", round(stats.worked_hours, 2)),
            ("Отработано дней", stats.worked_days),
            ("Опозданий", stats.late_arrivals),
            ("Всего минут опоздания", stats.total_late_minutes),
            ("Ранних уходов", stats.early_departures),
            ("Всего минут раннего ухода", stats.total_early_departure_minutes),
            ("Дней без верхней формы", stats.days_without_upper_uniform),
            ("Дней без нижней формы", stats.days_without_lower_uniform),
            ("Дней болезни", stats.sick_days),
            ("Дней отпуска", stats.vacation_days),
            ("Неуважительных отсутствий", stats.unexcused_absences),
            ("Незаполненных дней", stats.not_filled_days),
        ]

    # --------------------------------------------------------------- Excel

    async def export_excel(self, coach_id: int, year: int, month: int) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        coach, rows, stats = await self._gather(coach_id, year, month)

        wb = Workbook()
        ws = wb.active
        ws.title = "Отчёт"

        header_fill = PatternFill("solid", start_color="366092")
        header_font = Font(bold=True, color="FFFFFF")
        title_font = Font(bold=True, size=14)
        thin = Side(style="thin", color="BFBFBF")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        missing_fill = PatternFill("solid", start_color="FFF2CC")

        ws["A1"] = "Ежемесячный отчёт о работе тренера"
        ws["A1"].font = title_font
        ws.merge_cells("A1:J1")

        meta = [
            ("Тренер", coach.full_name),
            ("Табельный номер", coach.unique_id),
            ("Должность", coach.position),
            ("Команда", coach.team),
            ("Отделение", coach.branch.name if coach.branch else "—"),
            ("Период", f"{MONTHS_RU[month]} {year}"),
        ]
        for offset, (label, value) in enumerate(meta, start=3):
            ws.cell(row=offset, column=1, value=label).font = Font(bold=True)
            ws.cell(row=offset, column=2, value=value)

        headers = [
            "Дата", "Присутствие", "Верх. форма", "Нижн. форма",
            "Начало", "Окончание", "Часов", "Опоздание, мин",
            "Ранний уход, мин", "Примечания",
        ]
        header_row = 3 + len(meta) + 1
        for col, title in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col, value=title)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        row_idx = header_row + 1
        for report in rows:
            values = [
                report.report_date.strftime("%d.%m.%Y"),
                ATTENDANCE_RU.get(report.attendance, "—"),
                UNIFORM_RU.get(report.upper_uniform, "—"),
                UNIFORM_RU.get(report.lower_uniform, "—"),
                report.start_time.strftime("%H:%M") if report.start_time else "—",
                report.end_time.strftime("%H:%M") if report.end_time else "—",
                float(report.worked_hours or 0),
                report.late_arrival_minutes,
                report.early_departure_minutes,
                report.notes or "",
            ]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.border = border
                # Highlight days nobody filled in, so gaps are visible at a glance.
                if not report.is_completed:
                    cell.fill = missing_fill
            row_idx += 1

        row_idx += 1
        ws.cell(row=row_idx, column=1, value="Итоговая статистика").font = Font(
            bold=True, size=12
        )
        row_idx += 1
        for label, value in self._summary_rows(stats):
            ws.cell(row=row_idx, column=1, value=label).font = Font(bold=True)
            ws.cell(row=row_idx, column=2, value=value)
            row_idx += 1

        widths = [12, 16, 14, 14, 10, 12, 9, 16, 18, 30]
        for col, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

        buffer = io.BytesIO()
        wb.save(buffer)
        logger.info("Excel export ready for coach={} {}/{}", coach_id, month, year)
        return buffer.getvalue()

    # ----------------------------------------------------------------- PDF

    async def export_pdf(self, coach_id: int, year: int, month: int) -> bytes:
        if not _register_fonts():
            raise RuntimeError(
                "PDF export requires a Cyrillic font (fonts-dejavu-core)"
            )

        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
        )

        coach, rows, stats = await self._gather(coach_id, year, month)

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=A4,
            topMargin=15 * mm, bottomMargin=15 * mm,
            leftMargin=12 * mm, rightMargin=12 * mm,
            title=f"{coach.full_name} - {MONTHS_RU[month]} {year}",
        )

        base = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "TitleRu", parent=base["Heading1"], fontName=FONT_BOLD,
            fontSize=15, textColor=colors.HexColor("#366092"), spaceAfter=6,
        )
        body_style = ParagraphStyle(
            "BodyRu", parent=base["Normal"], fontName=FONT_REGULAR, fontSize=9
        )

        story: list = [
            Paragraph("Ежемесячный отчёт о работе тренера", title_style),
            Paragraph(
                f"{coach.full_name} — {MONTHS_RU[month]} {year}", body_style
            ),
            Spacer(1, 6 * mm),
        ]

        meta = [
            ["Табельный номер", coach.unique_id],
            ["Должность", coach.position],
            ["Команда", coach.team],
            ["Отделение", coach.branch.name if coach.branch else "—"],
        ]
        meta_table = Table(meta, colWidths=[45 * mm, 100 * mm])
        meta_table.setStyle(
            TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), FONT_REGULAR),
                ("FONTNAME", (0, 0), (0, -1), FONT_BOLD),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ])
        )
        story.extend([meta_table, Spacer(1, 6 * mm)])

        table_data: list[list[str]] = [[
            "Дата", "Присутствие", "Верх", "Низ", "Начало",
            "Конец", "Часов", "Опозд.", "Ранний уход",
        ]]
        incomplete_rows: list[int] = []
        for position, report in enumerate(rows, start=1):
            if not report.is_completed:
                incomplete_rows.append(position)
            table_data.append([
                report.report_date.strftime("%d.%m"),
                ATTENDANCE_RU.get(report.attendance, "—"),
                UNIFORM_RU.get(report.upper_uniform, "—"),
                UNIFORM_RU.get(report.lower_uniform, "—"),
                report.start_time.strftime("%H:%M") if report.start_time else "—",
                report.end_time.strftime("%H:%M") if report.end_time else "—",
                f"{float(report.worked_hours or 0):.1f}",
                str(report.late_arrival_minutes),
                str(report.early_departure_minutes),
            ])

        col_widths = [16, 28, 20, 20, 18, 18, 16, 18, 26]
        table = Table(
            table_data,
            colWidths=[w * mm for w in col_widths],
            repeatRows=1,
        )
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#366092")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), FONT_BOLD),
            ("FONTNAME", (0, 1), (-1, -1), FONT_REGULAR),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BFBFBF")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")]),
        ]
        for row_number in incomplete_rows:
            style.append(
                ("BACKGROUND", (0, row_number), (-1, row_number), colors.HexColor("#FFF2CC"))
            )
        table.setStyle(TableStyle(style))
        story.extend([table, Spacer(1, 6 * mm)])

        story.append(Paragraph("Итоговая статистика", title_style))
        summary = [[label, str(value)] for label, value in self._summary_rows(stats)]
        summary_table = Table(summary, colWidths=[70 * mm, 40 * mm])
        summary_table.setStyle(
            TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), FONT_REGULAR),
                ("FONTNAME", (0, 0), (0, -1), FONT_BOLD),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BFBFBF")),
                ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")]),
            ])
        )
        story.append(summary_table)

        doc.build(story)
        logger.info("PDF export ready for coach={} {}/{}", coach_id, month, year)
        return buffer.getvalue()
